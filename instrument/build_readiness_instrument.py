from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'Task_Conditioned_Data_Readiness_Review_2026-10-23.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.62), Inches(.62)
sec.left_margin, sec.right_margin = Inches(.7), Inches(.7)
sec.footer_distance = Inches(.28)
styles = doc.styles
for name in ['Normal','Title','Heading 1','Heading 2']:
    st = styles[name]
    st.font.name = 'Times New Roman'
    st.font.color.rgb = RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
styles['Normal'].font.size = Pt(11)
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Normal'].paragraph_format.line_spacing = 1.02
styles['Title'].font.size = Pt(17)
styles['Title'].paragraph_format.space_after = Pt(4)
styles['Title'].paragraph_format.space_before = Pt(0)
styles['Heading 1'].font.size = Pt(12)
styles['Heading 1'].paragraph_format.space_before = Pt(8)
styles['Heading 1'].paragraph_format.space_after = Pt(5)
styles['Heading 1'].paragraph_format.keep_with_next = True
for stname in ['Title','Heading 1','Heading 2']:
    pp=styles[stname]._element.find(qn('w:pPr'))
    if pp is not None:
        for el in list(pp):
            if el.tag == qn('w:pBdr'): pp.remove(el)

plain=[]
counter=10
def p(text='',style=None,boldlead=None,after=None):
    pp=doc.add_paragraph(style=style)
    if boldlead:
        pp.add_run(boldlead).bold=True
        pp.add_run(text)
    else: pp.add_run(text)
    if after is not None: pp.paragraph_format.space_after=Pt(after)
    plain.append((boldlead or '')+text)
    return pp

def sdt(pp,placeholder):
    global counter
    counter+=1
    el=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr')
    alias=OxmlElement('w:alias');alias.set(qn('w:val'),placeholder)
    eid=OxmlElement('w:id');eid.set(qn('w:val'),str(counter))
    tx=OxmlElement('w:text');pr.extend([alias,eid,tx]);el.append(pr)
    content=OxmlElement('w:sdtContent');r=OxmlElement('w:r')
    t=OxmlElement('w:t');t.text='['+placeholder+']';r.append(t);content.append(r)
    el.append(content);pp._p.append(el)

def fields(items,after=3):
    pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(after)
    chunks=[]
    for i,(label,placeholder) in enumerate(items):
        if i: pp.add_run('    ')
        pp.add_run(label+' ').italic=True;sdt(pp,placeholder)
        chunks.append(label+' ['+placeholder+']')
    plain.append('    '.join(chunks))
    return pp

def criterion(label,text):
    pp=p(text,boldlead=label+' ',after=2)
    pp.paragraph_format.keep_with_next=True
    q=fields([('Documentation','0 / 1 / 2 / U / N/A'),('Adequacy','S / C / N / U')],after=1)
    q.paragraph_format.keep_with_next=True
    q=fields([('Evidence','file page measurement or link'),('Action and owner','enter')],after=6)
    # Keep the two response lines attached to their criterion, but allow a
    # complete criterion to move as a unit if a response expands in Word.
    q.paragraph_format.keep_with_next=False

p('Task conditioned data readiness review','Title')
p('23 October 2026  |  Research instrument for author review',after=6)
p('Complete one profile for a holding, version and intended use. Analysis readiness makes data interpretable and usable; AI readiness adds evidence about learning and the claimed generalization. This instrument will be refined through research. It is not an NSF requirement or a validated universal standard.')
p('Declare the use case','Heading 1')
fields([('Holding and DOI or version','enter'),('Reviewer and date','enter')])
fields([('Scientific question and intended user','enter')])
fields([('Inputs and their roles','enter'),('Output or learning objective','enter')])
fields([('Learning regime','supervised / self-supervised / semi-supervised / pseudo-label self-training / other')])
fields([('Independent unit and native spatial or temporal support','enter')])
fields([('Intended use and setting','local inference new region physical regime cloud or edge')])
fields([('Acceptance criteria failure costs and uncertainty requirements','enter')],after=7)
p('For each criterion, record evidence and the next action. Documentation: 0 = a required element is confirmed missing; 1 = incomplete; 2 = another team can verify it; U = unassessed or insufficient evidence; N/A = inapplicable with a reason. Adequacy for this use: S = supported; C = conditional; N = not supported; U = unresolved. A documented limitation can score 2 and still make a use unsupported. Do not sum grades or impose a universal pass mark.',after=6)
p('Analysis prerequisites','Heading 1')
criterion('A1 Provenance and permission','Identify source, version or retrieval date, citation, permitted access, learning use and redistribution. Trace derived records to processing and environment information. Rights or provenance that remain unknown must stay visible.')
criterion('A2 Homogenization and physical comparability','Check units, time base, reference frame, response and instrument changes. Preserve native support and uncertainty; document transformations and observed, simulated or derived origin. A common shape does not establish equivalent measurements.')
criterion('A3 Quality and gaps','Document quality flags, gaps, sentinels, censoring, detection limits, exclusions and their causes. Determine whether cleaning or selection changes the population or removes information needed for this use.')
criterion('A4 Structure and metadata','Verify readable variables, axes, units, relationships and sample manifests, plus a reproducible loader or partial read. Metadata conformance and object storage support access; they do not alone establish scientific fitness.')

doc.add_page_break();plain.append('\f')
p('Learning and generalization','Heading 1')
p('Apply these criteria to the declared objective and model or learning method. Record measured quantities where possible. Unknown evidence remains U; explain any N/A. Preserve the scientific distinctions between measurement families, representations and source origins.',after=7)
criterion('B1 Diversity and independence','Count independent events, sites, regions, years and simulation families alongside rows and bytes. Document dependence and exclusions. Assess coverage against the intended population or physical design space, rather than the archive alone.')
criterion('B2 Distribution','Characterize inputs, targets where present, rare conditions, imbalance and missing modalities. State the reference population, sampling policy and known shifts. Separate natural frequencies from deliberate training reweighting.')
criterion('B3 Noise','Distinguish instrument and processing uncertainty from physical variability. State which variability is signal for this question. Check whether filtering, compression or normalization suppresses information needed for the intended inference.')
criterion('B4 Augmentation','Justify transformations, masking or perturbations using measurement physics and task semantics. Test preservation of relevant observables and targets. Validity depends on the sensor, phenomenon and question; standard augmentation recipes are not universally applicable.')
criterion('B5 Labels or learning structure','Supervised tasks need defined targets, label provenance, consistency, uncertainty and suitable negatives. Self-supervised objectives need meaningful context, views, pairing and sampling; external training labels are optional, independent evaluation is not. Pseudo-label self-training requires separate teacher provenance and error checks.')
criterion('B6 Dimensionality','State sample axes, shape, native resolution and independent samples relative to features. Test whether reduction retains task information, including rare behavior. Record bytes read, throughput, memory and compute at the tested scale; explained variance alone is insufficient.')
criterion('B7 Partitioning','Use independent groups and protect final tests from tuning, preprocessing and related-record leakage. Match tests to claims: statistical distance from training support; spatiotemporal transfer to new places or times; physical transfer to new event size, process or physics. Report within-domain and transfer results separately, with uncertainty and failures.')

doc.add_page_break();plain.append('\f')
p('Use setting and review decision','Heading 1')
p('Deployment fit is conditional on the declared research workflow. Record the evidence below for cloud, laboratory or sensing-edge use, as applicable. Research demonstrations do not establish operational-service reliability.',after=7)
fields([('Access route permissions and available inputs','enter')],after=7)
fields([('Measured throughput memory latency and transfer or compute cost','enter')],after=7)
fields([('Instrument interfaces missing inputs and known shifts','enter')],after=7)
fields([('Monitoring and observation prediction discrepancy checks','enter')],after=7)
fields([('Response to mismatch','retrieve evidence qualify abstain or revise after review')],after=7)
fields([('Evidence source next action owner and review date','enter')],after=7)
p('Community maturity is a separate profile. Record maintained tools and releases, stewardship, support and independent reuse or replication, including dates and scope. Independently checkable documentation does not prove independent use, and adoption does not establish fitness for a task.',after=7)
fields([('Stewardship maintenance support and contribution route','enter')],after=7)
fields([('Independent use or replication and its evidence','enter')],after=7)
p('Review conclusion','Heading 1')
fields([('Supported use and boundary of the evidence','enter')],after=7)
fields([('Conditional or unsupported uses and why','enter')],after=7)
fields([('Unresolved evidence or prerequisites','enter')],after=7)
fields([('Priority actions owners and due dates','enter')],after=7)
fields([('Reviewer or steward confirmation','enter')],after=7)
p('Reassess when the dataset version, model, objective or intended use changes. An analysis-ready station record can support a local study while remaining insufficient to test transfer to another region. Unlabeled continuous DAS can support a self-supervised objective; supervised phase picking additionally needs trustworthy targets and task-appropriate evaluation. Neither example receives a readiness judgment without evidence.',after=8)
p('Basis and interpretation','Heading 1')
p('Adapted from Marine Denolle, MLGeo Session 10, AI-ready data, 23 October 2026, Book section 2.13. The lecture distinguishes A1–A4 analysis prerequisites and B1–B7 learning criteria. Its 0–2 course scale concerns documentation; this research adaptation adds explicit U and N/A, a separate task-adequacy judgment and the use-setting review.',after=4)
p('https://geo-smart.github.io/mlgeo-book/slides/2026/lec10_ai_ready_data.html',after=0)

foot=sec.footer.paragraphs[0]
foot.alignment=2
r=foot.add_run('Page ');r.font.size=Pt(9)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
for r in foot.runs:r.font.name='Times New Roman';r.font.color.rgb=RGBColor(0,0,0)

doc.core_properties.title='Task conditioned data readiness review'
doc.core_properties.subject='Task-specific data readiness author review instrument'
doc.core_properties.author='Marine Denolle'
doc.core_properties.keywords='AI readiness, analysis readiness, geoscience, research instrument'
doc.save(OUT)
txt=OUT.with_suffix('.txt');txt.write_text('\n\n'.join(plain)+'\n')
print(json.dumps({'docx':str(OUT),'txt':str(txt),'words':len(' '.join(plain).split()),'fillable_controls':counter-10},indent=2))
